import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DATA_DIR = Path("data/tags")
OUT = DATA_DIR / "review_tag_dataset.xlsx"

SHEETS = [
    ("summary", "summary.csv"),
    ("tag_dictionary", "review_tag_dictionary.csv"),
    ("keyword_counts", "review_keyword_counts.csv"),
    ("extracted_headings", "review_extracted_issue_headings.csv"),
    ("header_dataset", "review_header_dataset.csv"),
    ("header_dataset_expanded", "review_header_dataset_expanded.csv"),
    ("unreadable_files", "unreadable_files.csv"),
]


def read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    return rows[0], rows[1:]


def add_sheet(wb: Workbook, title: str, csv_name: str) -> None:
    ws = wb.create_sheet(title=title[:31])
    headers, rows = read_csv(DATA_DIR / csv_name)
    ws.append(headers)
    for row in rows:
        ws.append(row)

    fill = PatternFill("solid", fgColor="176B87")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows[:500]:
            if idx <= len(row):
                max_len = max(max_len, min(len(str(row[idx - 1])), 80))
        width = min(max(max_len + 2, 10), 48)
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for title, csv_name in SHEETS:
        add_sheet(wb, title, csv_name)
    wb.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
